"""Excel export carries the caveats a CSV cannot (FR-024, §17).

"99.2% available" in a management pack, with nobody knowing that maintenance
windows were excluded and one application had no data for two days, is worse
than no figure. So the workbook states what it means alongside what it says.
"""
from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.services.report_service import export_availability_xlsx


ROWS = [
    {"application_name": "HRMS", "environment": "Production", "total_checks": 100,
     "successful_checks": 99, "availability_percent": 99.0, "sla_target_percent": 99.5,
     "sla_met": False, "avg_response_time": 812.4, "incident_count": 1,
     "avg_downtime_seconds": 300, "total_downtime_seconds": 300},
    {"application_name": "ITSM", "environment": "Production", "total_checks": 100,
     "successful_checks": 100, "availability_percent": 100.0, "sla_target_percent": None,
     "sla_met": None, "avg_response_time": 120.0, "incident_count": 0,
     "avg_downtime_seconds": 0, "total_downtime_seconds": 0},
]


def _book(meta=None):
    return load_workbook(BytesIO(export_availability_xlsx(ROWS, meta)))


def test_the_figures_sheet_holds_every_row(db):
    sheet = _book()["Availability"]
    assert sheet.max_row == 3          # header plus two applications
    assert sheet["A2"].value == "HRMS"
    assert sheet["E2"].value == 99.0


def test_a_missing_sla_target_says_so_rather_than_showing_a_number(db):
    """An invented 99.9 would mark an application as failing something nobody
    asked it to meet."""
    sheet = _book()["Availability"]
    assert sheet["F3"].value == "no target"
    assert sheet["G3"].value == "-"
    assert sheet["G2"].value == "missed"


def test_the_workbook_explains_how_the_figures_were_calculated(db):
    notes = _book({"From": "2026-08-01", "To": "2026-08-31", "Timezone": "UTC"})["About this report"]
    text = " ".join(str(c.value) for row in notes.iter_rows() for c in row if c.value)
    assert "2026-08-01" in text and "UTC" in text
    assert "successful checks / total checks" in text
    assert "open incidents are excluded" in text
    assert "Data gaps" in text


def test_the_header_row_exists_and_is_frozen(db):
    """A spreadsheet scrolled past its header is a grid of unlabelled numbers."""
    sheet = _book()["Availability"]
    assert sheet["A1"].value == "Application"
    assert sheet.freeze_panes == "A2"
